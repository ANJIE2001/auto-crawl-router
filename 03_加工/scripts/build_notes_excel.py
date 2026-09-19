"""Build notes Excel from profile JSONs or browser data, matching the 24-field template.

CLI 用法 — 灵造路径（推荐，省 WB）:
  python build_notes_excel.py \
    --profile "布布糕" "布布糕.json" \
    --output "笔记原始数据表_2026-08-04.xlsx"

CLI 用法 — 浏览器路径（灵造不可用时的 fallback）:
  python build_notes_excel.py \
    --browser-data "浏览器原始数据/2026-08-04_profile_布布糕.json" \
    --output "笔记原始数据表_2026-08-04.xlsx"

v5: 双路径支持 — --profile (灵造) + --browser-data (浏览器 fallback)。
"""

import argparse
import json
import os
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

DEFAULT_BASE = Path(r"C:\Users\PC\WorkBuddy\edge浏览器查询")

HEADERS = [
    "序号", "博主名称", "博主ID", "标题", "笔记ID", "发布时间（UTC+8）", "类型", "时长（秒）",
    "点赞", "收藏", "评论", "分享", "互动总量", "收藏/点赞",
    "协作标记", "商品笔记", "原始标签/描述", "字幕全文索引",
    "笔记链接", "封面原始URL", "字幕状态", "xhs_note_type", "置顶", "数据来源",
]


def extract_notes(profile_name: str, json_path: str) -> list:
    """Extract all notes from a profile JSON, return list of dicts."""
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    items = data["data"]["items"]
    artifacts = data["data"]["artifacts"]
    subtitle_artifact_url = artifacts.get("subtitle_markdown", {}).get("url", "")

    user_data = data["data"].get("user", {})
    blogger_name = user_data.get("nickname", profile_name)
    blogger_id = user_data.get("id", "")

    notes = []
    for item in items:
        note_id = item.get("id", "")
        title = item.get("title", "")
        note_type = item.get("type", "")
        xhs_type = item.get("xhs_note_type", "")
        url = item.get("url", "")

        pub_str = item.get("published_at", "")
        pub_dt = None
        if pub_str:
            try:
                pub_dt = datetime.fromisoformat(pub_str.replace("Z", "+00:00"))
                pub_dt = pub_dt.astimezone(timezone(timedelta(hours=8)))
                pub_dt = pub_dt.replace(tzinfo=None)
            except (ValueError, TypeError):
                pass

        sticky = item.get("sticky", False)
        media = item.get("media", {})
        duration = media.get("video_duration_seconds")
        cover_url = media.get("cover_large_url", "")

        metrics = item.get("metrics", {})
        liked = metrics.get("liked", 0) or 0
        shared = metrics.get("shared", 0) or 0
        collected = metrics.get("collected", 0) or 0
        commented = metrics.get("commented", 0) or 0
        total_interaction = liked + shared + collected + commented
        collect_like_ratio = round(collected / liked, 4) if liked > 0 else 0

        monet = item.get("monetization", {})
        collab_tag = "是" if monet.get("collaboration", {}).get("likely_collaboration", False) else "否"
        commerce_note = "是" if monet.get("commerce_note", {}).get("is_goods_note", False) else "否"

        text = item.get("text", {})
        tags_desc = text.get("desc", "")
        subtitle = text.get("subtitle", {})
        subtitle_status = subtitle.get("status", "")
        subtitle_index_url = subtitle_artifact_url if subtitle_status == "ready" else ""

        notes.append({
            "博主名称": blogger_name,
            "博主ID": blogger_id,
            "标题": title,
            "笔记ID": note_id,
            "发布时间": pub_dt,
            "类型": note_type,
            "时长（秒）": duration,
            "点赞": liked,
            "收藏": collected,
            "评论": commented,
            "分享": shared,
            "互动总量": total_interaction,
            "收藏/点赞": collect_like_ratio,
            "协作标记": collab_tag,
            "商品笔记": commerce_note,
            "原始标签/描述": tags_desc,
            "字幕全文索引": subtitle_index_url,
            "笔记链接": url,
            "封面原始URL": cover_url,
            "字幕状态": subtitle_status,
            "xhs_note_type": xhs_type,
            "置顶": "是" if sticky else "否",
            "数据来源": f"analyze-user-profile",
        })
    return notes


def build_excel(all_notes: list, output_path: str, title_suffix: str = ""):
    wb = Workbook()
    ws = wb.active
    ws.title = "笔记原始数据表"

    header_font = Font(name="微软雅黑", size=11, bold=True)
    data_font = Font(name="微软雅黑", size=10)
    title_font = Font(name="微软雅黑", size=14, bold=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.cell(1, 1, f"笔记原始数据表｜{len(all_notes)} 条笔记{title_suffix}").font = title_font
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Row 2: separator
    ws.row_dimensions[2].height = 6

    # Row 3: Headers
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(3, col_idx, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = center_align
    ws.row_dimensions[3].height = 24

    # Data rows
    center_cols = {1, 2, 3, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 21, 22, 23, 24}
    for row_idx, note in enumerate(all_notes, 1):
        excel_row = row_idx + 3
        values = [
            row_idx,
            note["博主名称"], note["博主ID"], note["标题"], note["笔记ID"],
            note["发布时间"], note["类型"], note["时长（秒）"],
            note["点赞"], note["收藏"], note["评论"], note["分享"],
            note["互动总量"], note["收藏/点赞"],
            note["协作标记"], note["商品笔记"],
            note["原始标签/描述"], note["字幕全文索引"],
            note["笔记链接"], note["封面原始URL"],
            note["字幕状态"], note["xhs_note_type"],
            note["置顶"], note["数据来源"],
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(excel_row, col_idx, val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if col_idx in center_cols else left_align
            if isinstance(val, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if col_idx == 14 and isinstance(val, (int, float)):
                cell.number_format = "0.0000"
        ws.row_dimensions[excel_row].height = 22

    # Column widths
    widths = {1:5, 2:14, 3:24, 4:36, 5:22, 6:20, 7:7, 8:9, 9:8, 10:8, 11:7, 12:7,
              13:10, 14:10, 15:8, 16:8, 17:42, 18:50, 19:40, 20:50, 21:10, 22:13, 23:7, 24:30}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(HEADERS))}{3 + len(all_notes)}"

    wb.save(output_path)
    return output_path


def extract_notes_from_browser(json_path: str) -> list:
    """Extract notes from browser-normalized JSON (fallback path)."""
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    blogger_name = data.get("blogger_name", "unknown")
    blogger_id = data.get("blogger_id", "")
    notes_raw = data.get("notes", [])

    notes = []
    for note in notes_raw:
        note_id = note.get("noteId", "")
        title = note.get("title", "")
        note_type = note.get("noteType", "normal")
        publish_ts = note.get("publishTime", 0)

        pub_dt = None
        if publish_ts:
            try:
                pub_dt = datetime.fromtimestamp(publish_ts / 1000,
                                                tz=timezone(timedelta(hours=8)))
                pub_dt = pub_dt.replace(tzinfo=None)
            except (ValueError, OSError):
                pass

        liked = safe_int(note.get("likedCount", 0))
        collected = safe_int(note.get("collectedCount", 0))
        commented = safe_int(note.get("commentCount", 0))
        shared = safe_int(note.get("shareCount", 0))
        total_interaction = liked + shared + collected + commented
        collect_like_ratio = round(collected / liked, 4) if liked > 0 else 0

        duration = note.get("videoDuration", 0)
        cover_url = note.get("coverUrl", "")
        desc = note.get("desc", "")
        tags_list = note.get("tags", [])
        tags_str = " ".join(f"#{t}" for t in tags_list) if tags_list else desc
        note_url = note.get("noteUrl",
                            f"https://www.xiaohongshu.com/explore/{note_id}")
        is_top = note.get("isTop", False)
        subtitle_text = note.get("subtitleText", "")

        if subtitle_text:
            subtitle_status = "browser_extracted"
        elif note_type == "video":
            subtitle_status = "browser_no_subtitle"
        else:
            subtitle_status = "not_applicable"

        notes.append({
            "博主名称": blogger_name,
            "博主ID": blogger_id,
            "标题": title,
            "笔记ID": note_id,
            "发布时间": pub_dt,
            "类型": note_type,
            "时长（秒）": duration,
            "点赞": liked,
            "收藏": collected,
            "评论": commented,
            "分享": shared,
            "互动总量": total_interaction,
            "收藏/点赞": collect_like_ratio,
            "协作标记": "未知",
            "商品笔记": "未知",
            "原始标签/描述": tags_str,
            "字幕全文索引": "",
            "笔记链接": note_url,
            "封面原始URL": cover_url,
            "字幕状态": subtitle_status,
            "xhs_note_type": note_type,
            "置顶": "是" if is_top else "否",
            "数据来源": "浏览器CDP",
        })
    return notes


def safe_int(val) -> int:
    """兼容字符串数字和整数"""
    if val is None:
        return 0
    try:
        return int(val)
    except (ValueError, TypeError):
        return 0


def main():
    parser = argparse.ArgumentParser(description="从 profile JSON 生成笔记数据 Excel")
    parser.add_argument("--base-dir", default=str(DEFAULT_BASE),
                        help=f"项目根目录（默认: {DEFAULT_BASE}）")
    parser.add_argument("--json-dir", default=None,
                        help="灵造原始 JSON 目录（默认: BASE/内容资产库/90_数据源/灵造原始JSON）")

    # --profile "昵称" "路径" 可多次指定（灵造路径）
    parser.add_argument("--profile", nargs=2, action="append", default=[],
                        metavar=("NAME", "JSON_PATH"),
                        help="博主名称和 profile JSON 路径（可多次指定）")

    # --browser-data "路径" 可多次指定（浏览器 fallback 路径）
    parser.add_argument("--browser-data", action="append", default=[],
                        metavar="JSON_PATH",
                        help="浏览器标准化 JSON 路径（可多次指定）")

    parser.add_argument("--output", required=True,
                        help="输出 Excel 文件路径")
    parser.add_argument("--title-suffix", default="",
                        help="表格标题后缀（如日期）")

    args = parser.parse_args()

    if not args.profile and not args.browser_data:
        print("❌ 请提供 --profile 或 --browser-data", file=sys.stderr)
        sys.exit(1)

    base = Path(args.base_dir)
    json_dir = args.json_dir or str(base / "内容资产库" / "90_数据源" / "灵造原始JSON")

    all_notes = []

    # 灵造路径
    for name, json_path in args.profile:
        if os.path.isabs(json_path):
            json_abs = json_path
        else:
            candidate = os.path.join(json_dir, json_path)
            if os.path.exists(candidate):
                json_abs = candidate
            else:
                candidate2 = os.path.join(str(base), json_path)
                if os.path.exists(candidate2):
                    json_abs = candidate2
                else:
                    print(f"❌ 找不到文件: {json_path}", file=sys.stderr)
                    sys.exit(1)
        notes = extract_notes(name, json_abs)
        all_notes.extend(notes)
        print(f"  [灵造] {name}: {len(notes)} 条")

    # 浏览器 fallback 路径
    for json_path in args.browser_data:
        if os.path.isabs(json_path):
            json_abs = json_path
        else:
            candidate = os.path.join(str(base), json_path)
            if os.path.exists(candidate):
                json_abs = candidate
            else:
                print(f"❌ 找不到文件: {json_path}", file=sys.stderr)
                sys.exit(1)
        notes = extract_notes_from_browser(json_abs)
        all_notes.extend(notes)
        print(f"  [浏览器] {notes[0]['博主名称'] if notes else '?'}: {len(notes)} 条")

    # Sort by published_at descending
    all_notes.sort(key=lambda n: n["发布时间"] or datetime.min, reverse=True)
    video_count = sum(1 for n in all_notes if n["类型"] == "video")
    image_count = sum(1 for n in all_notes if n["类型"] == "normal")

    print(f"  总计: {len(all_notes)} 条 | 视频: {video_count} | 图文: {image_count}")
    print(f"  字幕ready: {sum(1 for n in all_notes if n['字幕状态'] == 'ready')}")

    # 输出路径处理同上
    if os.path.isabs(args.output):
        output = args.output
    else:
        output = str(base / args.output)
    build_excel(all_notes, output, args.title_suffix)
    print(f"  输出: {output}")


if __name__ == "__main__":
    main()
