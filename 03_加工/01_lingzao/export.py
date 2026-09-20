# -*- coding: utf-8 -*-
"""灵造的出表逻辑（Excel）。

版式：第 1 行大标题、第 2 行留空、第 3 行表头、第 4 行起数据。
表头用**蓝底**（`D9E1F2`）—— 和得到大脑的绿底区分开，一眼看出是哪个源。
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Iterable

from record import FIELD_LABELS, NoteRecord

# 列宽（按 FIELD_LABELS 的顺序给，不够的用默认）
COL_WIDTHS = {
    "序号": 5, "数据来源": 9, "采集命令": 17, "抓取时间": 19, "账号名称": 16,
    "平台账号名": 16, "账号ID": 24, "标题": 38, "笔记ID": 22, "发布时间(UTC+8)": 19,
    "类型": 8, "xhs_note_type": 13, "时长(秒)": 9, "点赞": 9, "收藏": 9,
    "评论": 8, "分享": 8, "互动总量": 10, "收藏/点赞": 10, "协作标记": 9,
    "商品笔记": 9, "置顶": 6, "原始标签/描述": 44, "标签": 24, "字幕状态": 10,
    "字幕是否截断": 11, "逐字稿": 60, "笔记链接": 40, "封面URL": 44, "原始文件": 30,
}


def build_excel(records: list[NoteRecord], out_path: Path, title: str = "") -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("缺 openpyxl。装一下：pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "统一数据"
    labels = FIELD_LABELS
    ncol = len(labels)

    title_font = Font(name="微软雅黑", size=14, bold=True)
    header_font = Font(name="微软雅黑", size=11, bold=True)
    data_font = Font(name="微软雅黑", size=10)
    border = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"))
    header_fill = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")   # 灵造：蓝底
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 第 1 行标题
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.cell(1, 1, title or f"统一数据表｜{len(records)} 条").font = title_font
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 6

    # 第 3 行表头
    for c, name in enumerate(labels, 1):
        cell = ws.cell(3, c, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center
    ws.row_dimensions[3].height = 26

    # 数据
    for i, rec in enumerate(records, 1):
        row = i + 3
        values = rec.to_row(i)
        for c, name in enumerate(labels, 1):
            cell = ws.cell(row, c, values[name])
            cell.font = data_font
            cell.border = border
            cell.alignment = left if name in ("标题", "原始标签/描述", "标签", "逐字稿",
                                              "笔记链接", "封面URL", "原始文件") else center
            if name == "收藏/点赞" and isinstance(values[name], (int, float)):
                cell.number_format = "0.0000"
        # 带逐字稿的行撑高一点，不然一格两千字全挤成一行
        ws.row_dimensions[row].height = 90 if values.get("逐字稿") else 22

    for c, name in enumerate(labels, 1):
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS.get(name, 14)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{3 + len(records)}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
