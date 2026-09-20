# -*- coding: utf-8 -*-
"""读「笔记原始数据表」这种成品 xlsx → 统一字段。

**为什么需要这个入口**：历史数据不一定留得下原始 JSON，可能只剩一张当时导出的表。
好在表里字段是全的（连「字幕全文索引」和「封面原始URL」都还留着），够还原。

两个坑：

1. **表头不在第 1 行。** 版式是

       第 1 行   合并标题（笔记原始数据表｜…）
       第 2 行   空行
       第 3 行   真表头
       第 4 行起 数据

   所以先扫出含「标题」「笔记ID」的那一行当真表头，别写死行号。

2. **按列名取值，不按列序号。** 老表 24 列、现在的表 29 列，按位置取早晚错位。
   列名还得归一化 —— `发布时间（UTC+8）` 和 `发布时间(UTC+8)` 全角半角都出现过，
   `博主名称` / `账号名称`、`封面URL` / `封面原始URL` 也一样。
"""

from __future__ import annotations

import re
from pathlib import Path

from record import AdapterResult, BaseAdapter, NoteRecord
from util import extract_tags, safe_int

SOURCE = "lingzao"


def _norm(s) -> str:
    """列名归一化：去掉空白和各类括号，全角半角一视同仁。"""
    return re.sub(r"[\s（）()【】\[\]]", "", str(s or "")).strip()


def _bool3(v):
    """是/否/未知 三态。认不出来就是 None，别瞎猜。"""
    t = _norm(v)
    if t in ("是", "true", "True", "1"):
        return True
    if t in ("否", "false", "False", "0"):
        return False
    return None


def load_rows(path: str | Path, scan: int = 15) -> tuple[list[dict], dict]:
    """把表读成 [{归一化列名: 值}, …]，顺带返回表头在第几行。"""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("读 xlsx 需要 openpyxl，装一下：pip install openpyxl") from e

    ws = load_workbook(path, read_only=True).worksheets[0]
    raw = [list(r) for r in ws.iter_rows(values_only=True)]

    header_at = -1
    for i, row in enumerate(raw[:scan]):
        cells = {_norm(c) for c in row if c is not None}
        if "标题" in cells or "笔记ID" in cells:
            header_at = i
            break
    if header_at < 0:
        raise ValueError(
            f"{Path(path).name} 前 {scan} 行里找不到表头（没看到「标题」列）。"
            "这不是一张笔记表，或者版式又变了。"
        )

    header = [_norm(c) for c in raw[header_at]]
    rows: list[dict] = []
    for row in raw[header_at + 1:]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        rows.append(dict(zip(header, row)))

    info = {"header_row": header_at + 1, "columns": [h for h in header if h],
            "sheet": ws.title}
    return rows, info


class TableAdapter(BaseAdapter):
    """成品表 → 统一字段。已导出的表再导一次，方便跟 JSON 来的记录并轨。"""

    source = SOURCE
    command = "table"
    suffixes = (".xlsx",)

    def parse(self, path: str | Path) -> AdapterResult:
        rows, info = load_rows(path)

        records: list[NoteRecord] = []
        collection_urls: set[str] = set()

        for row in rows:
            cells = {k: ("" if v is None else str(v).strip()) for k, v in row.items()}

            def pick(*names, _c=cells):
                for n in names:
                    v = _c.get(_norm(n), "")
                    if v:
                        return v
                return ""

            desc = pick("原始标签/描述", "正文", "描述")
            tag_str = pick("标签")
            tags = [t for t in tag_str.split() if t] if tag_str else extract_tags(desc)

            url_index = pick("字幕全文索引", "字幕索引", "完整字幕")
            if url_index:
                collection_urls.add(url_index)

            rec = NoteRecord(
                source=SOURCE,
                command=pick("采集命令") or pick("数据来源") or "analyze-user-profile",
                raw_file=str(path),
                collected_at=pick("抓取时间"),
                note_id=pick("笔记ID"),
                title=pick("标题"),
                author_id=pick("博主ID", "账号ID"),
                author_name=pick("博主名称", "账号名称"),
                url=pick("笔记链接"),
                note_type=pick("类型"),
                xhs_note_type=pick("xhs_note_type", "xhsnotetype"),
                published_at=pick("发布时间（UTC+8）", "发布时间(UTC+8)", "发布时间"),
                duration_seconds=safe_int(pick("时长（秒）", "时长(秒)")),
                liked=safe_int(pick("点赞")),
                collected=safe_int(pick("收藏")),
                commented=safe_int(pick("评论")),
                shared=safe_int(pick("分享")),
                is_sticky=bool(_bool3(pick("置顶"))),
                is_collaboration=_bool3(pick("协作标记")),
                is_commerce_note=_bool3(pick("商品笔记")),
                desc=desc,
                tags=tags,
                cover_url=pick("封面URL", "封面原始URL", "封面url"),
                subtitle_status=pick("字幕状态"),
                subtitle_truncated=bool(_bool3(pick("字幕是否截断"))),
                subtitle_text=pick("逐字稿", "字幕正文", "字幕纯文本"),
                raw={k: v for k, v in row.items() if k},
            ).finalize()
            records.append(rec)

        meta = {
            "table_info": info,
            "collection_urls": sorted(collection_urls),
            "author_name": records[0].author_name if records else "",
            "author_id": records[0].author_id if records else "",
        }
        return AdapterResult(SOURCE, self.command, str(path), records, meta)


def parse_table(path: str | Path) -> AdapterResult:
    return TableAdapter().parse(path)
