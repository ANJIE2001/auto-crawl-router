# -*- coding: utf-8 -*-
"""得到大脑的出表逻辑。

Excel 版式和灵造那边保持一致（第 1 行大标题、第 2 行留空、第 3 行表头、第 4 行起数据），
这样两张表长得一样，人读起来不用切换脑子。
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Iterable

from record import FIELD_LABELS, NoteRecord

# 列宽。得到大脑这边「逐字稿」是主角，给得比灵造宽
COL_WIDTHS = {
    "序号": 5, "数据来源": 9, "采集命令": 18, "抓取时间": 19, "账号名称": 16,
    "平台账号名": 16, "账号ID": 12, "标题": 42, "笔记ID": 20, "发布时间(UTC+8)": 19,
    "类型": 8, "xhs_note_type": 13, "时长(秒)": 9, "点赞": 9, "收藏": 9,
    "评论": 8, "分享": 8, "互动总量": 10, "收藏/点赞": 10, "协作标记": 9,
    "商品笔记": 9, "置顶": 6, "原始标签/描述": 40, "标签": 26, "字幕状态": 10,
    "字幕是否截断": 12, "逐字稿": 70, "笔记链接": 40, "封面URL": 30, "原始文件": 40,
}


def build_excel(records: list[NoteRecord], out_path: Path, title: str = "") -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("缺 openpyxl。装一下：pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "统一数据"
    labels = FIELD_LABELS
    ncol = len(labels)

    title_font = Font(name="微软雅黑", size=14, bold=True)
    header_font = Font(name="微软雅黑", size=11, bold=True)
    data_font = Font(name="微软雅黑", size=10)
    border = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"), bottom=Side(style="thin"))
    header_fill = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")  # 得到大脑用绿底
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.cell(1, 1, title or f"得到大脑数据表｜{len(records)} 条").font = title_font
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 6

    for c, name in enumerate(labels, 1):
        cell = ws.cell(3, c, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center
    ws.row_dimensions[3].height = 26

    for i, rec in enumerate(records, 1):
        row = i + 3
        values = rec.to_row(i)
        for c, name in enumerate(labels, 1):
            cell = ws.cell(row, c, values[name])
            cell.font = data_font
            cell.border = border
            cell.alignment = left if name in ("标题", "原始标签/描述", "标签", "逐字稿",
                                              "笔记链接", "封面URL", "原始文件") else center
            if name == "收藏/点赞" and isinstance(values[name], (int, float)):
                cell.number_format = "0.0000"
        ws.row_dimensions[row].height = 90 if values.get("逐字稿") else 22

    for c, name in enumerate(labels, 1):
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS.get(name, 14)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{3 + len(records)}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
