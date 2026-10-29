# -*- coding: utf-8 -*-
"""出 xlsx：一张总表 + 每平台一张分表。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import record

# 每个平台一个浅底色，一眼能分开（跟另外两个源一样，用颜色区分子表）
PLAT_FILL = {
    "小红书": "FCE4EC",
    "抖音": "E8EAF6",
    "视频号": "E0F2F1",
    "公众号": "FFF8E1",
}

COL_WIDTH = {
    "序号": 6,
    "平台": 10,
    "账号名称": 24,
    "平台账号名": 20,
    "账号ID": 30,
    "粉丝数": 12,
    "简介": 46,
    "认证": 16,
    "头像URL": 30,
    "主页链接": 42,
    "关键词": 10,
    "抓取时间": 21,
    "原始文件": 34,
}

HEAD_FILL = PatternFill("solid", fgColor="F1EFE8")
HEAD_FONT = Font(bold=True, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_sheet(ws, rows, fill=None, note=None):
    ws.append(record.COLUMNS)
    for c in range(1, len(record.COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill("solid", fgColor=fill) if fill else HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    for r in rows:
        ws.append(r)

    for i, col in enumerate(record.COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTH.get(col, 16)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(cell.column == 7))
        # 粉丝数右对齐 + 千分位
        fc = row[record.COLUMNS.index("粉丝数")]
        if isinstance(fc.value, int):
            fc.number_format = "#,##0"
            fc.alignment = Alignment(horizontal="right", vertical="center")
    ws.freeze_panes = "A2"


def export_xlsx(recs: list, path: Path, title=None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # 总表
    ws = wb.active
    ws.title = "全部"
    _write_sheet(ws, record.finalize(recs))

    # 各平台分表
    for plat, fill in PLAT_FILL.items():
        sub = [r for r in recs if r.get("platform") == plat]
        if not sub:
            continue
        s = wb.create_sheet(plat)
        _write_sheet(s, record.finalize(sub), fill=fill)

    # 说明页
    info = wb.create_sheet("说明")
    lines = [
        ["数据来源", "JustOneAPI（api.justoneapi.com）"],
        ["生成时间", __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["记录条数", len(recs)],
        ["平台分布", "、".join(f"{p} {sum(1 for r in recs if r.get('platform') == p)} 条"
                          for p in PLAT_FILL if any(r.get("platform") == p for r in recs))],
        ["", ""],
        ["列名说明", ""],
        ["账号名称", "平台上的显示名"],
        ["平台账号名", "小红书号 / 抖音号 / 微信号；公众号是 wxid"],
        ["账号ID", "内部唯一 ID：小红书 24 位、抖音 uid、视频号 docID、公众号 gh_ 开头"],
        ["粉丝数", "「未知」= 该平台接口没给这个字段（视频号就是这种情况），不要当 0 用"],
        ["主页链接", "小红书/抖音是按 ID 拼的网页地址，视频号和公众号没有公开主页"],
        ["认证", "小红书只标「是否官方认证」；抖音是认证文案；视频号是博主类型"],
        ["", ""],
        ["价格提示", "每条记录都是一次付费调用换来的，源数据在 02_储存/03_justoneapi"],
    ]
    for row in lines:
        info.append(row)
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 86
    for row in info.iter_rows(min_row=1, max_row=info.max_row):
        if row[0].value:
            row[0].font = Font(bold=True, size=10)
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)
    return path
