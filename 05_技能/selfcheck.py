# -*- coding: utf-8 -*-
"""项目自检 —— 一次跑出所有「静默错误」。

    python 05_技能/selfcheck.py
    python 05_技能/selfcheck.py --quiet     # 只报错误和警告

**为什么要有它**：`.workbuddy/skills/` 里那份纠错清单有 30 条，
AI 读了不一定每条都照做。**能脚本化的检查就别靠人记** ——
脚本跑一次给确定的答案，比让 AI 记住 30 条可靠。

这个项目最危险的一类 bug 是「不报错但是错的」：
文件名少个下划线、字幕撞了截断上限、20 条被读成 40 条 —— 全都 exit 0。
所以这里的检查都盯着「静默」这两个字。

**八项检查**：
    1. 密钥      会进 git 的文件里有没有真 key
    2. 命名      落盘文件名合不合规（前导下划线、有没有重复堆积）
    3. 逐字稿    长度分布 —— 一堆一样的数字说明还在用截断版
    4. 封面      图池张数 vs 条目数
    5. 对账      产出目录数 vs 源数据条目数
    6. 结构      四层的文件齐不齐
    7. 配置      有没有残留旧路径 / 旧动作名
    8. 未知列    表里「未知」占比（提醒：别当 0 算）

首次 clone 下来跑（还没有数据）不会崩，只会提示「还没抓过数据」。
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
STORE = ROOT / "02_储存"
PRODUCE = ROOT / "04_产出"

if sys.platform == "win32":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

ERROR, WARN, INFO = "错误", "警告", "提示"
_found: list[tuple[str, str, str]] = []


def add(level: str, title: str, detail: str = "") -> None:
    _found.append((level, title, detail))


def rel(p) -> str:
    try:
        return str(Path(p).relative_to(ROOT))
    except Exception:
        return str(p)


# ---------------------------------------------------------------------------
# 1. 密钥
# ---------------------------------------------------------------------------

SECRETS = [
    (re.compile(r"gk_live_[A-Za-z0-9]{16,}"), "得到大脑 API Key"),
    (re.compile(r"cli_[a-f0-9]{16,}"), "得到大脑 Client ID"),
    (re.compile(r"\bsk-[A-Za-z0-9]{20,}"), "sk- 开头的密钥"),
    (re.compile(r"gh[pousr]_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{20,}"), "GitHub Token"),
    (re.compile(r"AKIA[0-9A-Z]{16}"), "AWS Access Key"),
]

# 这些文件本来就不进 git（.gitignore 挡着），命中是正常的
IGNORED_OK = {"config.local.json", ".env"}

SKIP_DIRS = {".git", "__pycache__", "node_modules", ".venv", "02_储存", "04_产出"}


def check_secrets() -> None:
    hits: list[tuple[str, str, str]] = []
    for p in walk_files(ROOT):
        if p.name.startswith(".tmp_"):
            continue
        # 无扩展名的也要扫 —— .env 就没有后缀，漏了它等于白扫
        name = p.name.lower()
        if (p.suffix.lower() not in (".py", ".json", ".md", ".txt", ".yaml", ".yml",
                                     ".example", ".cfg", ".ini")
                and name != ".env" and not name.startswith(".env.")):
            continue
        try:
            text = p.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            continue
        for rx, label in SECRETS:
            for m in rx.finditer(text):
                hits.append((rel(p), label, m.group(0)[:12] + "…"))

    real = [h for h in hits if Path(h[0]).name not in IGNORED_OK]
    benign = [h for h in hits if Path(h[0]).name in IGNORED_OK]

    if real:
        for f, label, val in real[:6]:
            add(ERROR, "进 git 的文件里发现了疑似密钥",
                f"{f}  →  {label}  {val}")
        if len(real) > 6:
            add(ERROR, "还有更多密钥命中", f"共 {len(real)} 处")
    if benign:
        add(INFO, "本地凭证文件存在（.gitignore 挡着，不进 git）",
            "、".join(sorted({Path(h[0]).name for h in benign}))
            + "　← 打包发人前也要排掉，见 05_技能/打包排除清单.md")
    if not hits:
        add(INFO, "密钥扫描",
            "干净，项目里没扫到真凭证 ✓（凭证该存本机，不在项目里）"
            "　→ 看当前存在哪：python 05_技能/set_key.py")


# ---------------------------------------------------------------------------
# 2. 命名
# ---------------------------------------------------------------------------

def check_naming() -> None:
    # 得到大脑：文件名必须带前导下划线（加工层靠 `*_blogger-content_*` 认）
    gn = STORE / "02_getnote" / "blogger"
    if gn.is_dir():
        bad = [f.name for f in gn.glob("*.json")
               if re.match(r"^(bloggers_|blogger-content|blogger-contents)", f.name)]
        if bad:
            add(ERROR, "得到大脑落盘文件名缺前导下划线 → 加工层认不出",
                f"{rel(gn)}  里 {len(bad)} 个，例：{bad[:3]}")
        dups = [f.name for f in gn.glob("*.json") if re.search(r"_\d+\.json$", f.name)
                and not f.name.startswith("2026") and not re.search(r"_\d{6,}\.json$", f.name)]
        if dups:
            add(ERROR, "得到大脑落了重复文件（那个源是「覆盖」语义，不该出现 _2）",
                f"{rel(gn)}：{dups[:3]}")
        if not bad and not dups:
            add(INFO, "得到大脑落盘命名", "合规")


# ---------------------------------------------------------------------------
# 3. 逐字稿
# ---------------------------------------------------------------------------

def check_subtitles() -> None:
    pack = PRODUCE / "博主"
    if not pack.is_dir():
        return
    for author_dir in sorted(p for p in pack.iterdir() if p.is_dir()):
        lens: list[int] = []
        empty = 0
        for d in author_dir.iterdir():
            if not d.is_dir():
                continue
            f = d / "逐字稿.md"
            if not f.is_file():
                continue
            t = f.read_text(encoding="utf-8", errors="ignore")
            body = t.split("## 逐字稿", 1)[-1].strip()
            if len(body) < 20 or "（这条没有逐字稿）" in body:
                empty += 1
            else:
                lens.append(len(body))
        if len(lens) >= 5:
            uniq = len(set(lens))
            if uniq <= 2:
                add(WARN, "逐字稿长度几乎一样 → 多半还在用截断版",
                    f"{author_dir.name}：{len(lens)} 条只要 {uniq} 种长度"
                    f"（撞同一个上限）。查合集有没有下到")
            else:
                add(INFO, f"逐字稿长度各异（{uniq} 种）说明是全文",
                    f"{author_dir.name}：有稿 {len(lens)} ｜ 无稿 {empty}")


# ---------------------------------------------------------------------------
# 4. 封面
# ---------------------------------------------------------------------------

def check_covers() -> None:
    pack = PRODUCE / "博主"
    pool = PRODUCE / "图片"
    if not pack.is_dir():
        return
    for author_dir in sorted(p for p in pack.iterdir() if p.is_dir()):
        dirs = [d for d in author_dir.iterdir() if d.is_dir()]
        if not dirs:
            continue
        have = sum(1 for d in dirs if any(f.name.startswith("封面") for f in d.iterdir()))
        pooled = 0
        if pool.is_dir():
            for pd in pool.iterdir():
                if pd.is_dir() and pd.name.endswith(author_dir.name.split("_")[-1]):
                    pooled = len([f for f in pd.iterdir() if f.is_file()])
        missing = len(dirs) - have
        if missing > len(dirs) * 0.3:
            add(WARN, "封面缺得多",
                f"{author_dir.name}：{have}/{len(dirs)}（图池 {pooled} 张）。"
                f"灵造的链接 3 小时过期，晚了补不回来")
        else:
            add(INFO, "封面覆盖",
                f"{author_dir.name}：{have}/{len(dirs)}（图池 {pooled} 张）")


# ---------------------------------------------------------------------------
# 5. 对账
# ---------------------------------------------------------------------------

def check_reconcile() -> None:
    pack = PRODUCE / "博主"

    # 得到大脑：详情 JSON 数 vs 产出目录数
    gn = STORE / "02_getnote" / "blogger"
    if gn.is_dir() and pack.is_dir():
        n_detail = len(list(gn.glob("*_blogger-content_*.json")))
        if n_detail:
            # 博主身份认 follow_id —— 产出目录名是「昵称_follow_id」。
            # ⚠️ 别按「条目数最接近」猜目录：两个博主都是 20 条时会认错人。
            fids = set()
            for f in gn.glob("*_bloggers_*.json"):
                try:
                    data = json.loads(f.read_text(encoding="utf-8")).get("data") or {}
                    for b in data.get("bloggers") or []:
                        if b.get("follow_id_str"):
                            fids.add(str(b["follow_id_str"]))
                except Exception:
                    pass
            hit = [d for d in pack.iterdir()
                   if d.is_dir() and any(d.name.endswith("_" + fid) for fid in fids)]
            for d in hit:
                n = len([x for x in d.iterdir() if x.is_dir()])
                if n < n_detail:
                    add(WARN, "有详情没归档到产出",
                        f"{d.name}：详情 {n_detail} 条，产出 {n} 个目录"
                        f" —— 跑一次 run.py blogger")
                else:
                    add(INFO, "得到大脑对账",
                        f"{d.name}：详情 {n_detail} ｜ 产出 {n}")

    # 灵造：profile 里的 items 数 vs 产出目录数
    lz = STORE / "01_lingzao" / "analyze-user-profile"
    if lz.is_dir() and pack.is_dir():
        for f in sorted(lz.glob("*.json")):
            try:
                j = json.loads(f.read_text(encoding="utf-8"))
            except Exception:
                continue
            data = j.get("data") or {}
            if data.get("type") != "analyze-user-profile":
                continue
            items = len(data.get("items") or [])
            user = (data.get("user") or {}).get("nickname") or "?"
            if not items:
                continue
            hit = [d for d in pack.iterdir()
                   if d.is_dir() and d.name.startswith(str(user)[:6])]
            if hit:
                n_pack = len([x for x in hit[0].iterdir() if x.is_dir()])
                tag = "✓" if n_pack >= items else "少了"
                add(INFO if n_pack >= items else WARN, f"灵造对账（{user}）",
                    f"{f.name}：源 {items} 条 ｜ 产出 {n_pack} 个目录  {tag}")


# ---------------------------------------------------------------------------
# 6. 结构
# ---------------------------------------------------------------------------

EXPECT_COLLECT = {"collect.py", "config.json", "README.md"}
EXPECT_PROCESS = {
    "01_lingzao": {"adapter.py", "bundle.py", "cover.py", "export.py",
                   "record.py", "run.py", "table.py", "util.py", "video.py"},
    "02_getnote": {"adapter.py", "bundle.py", "cover.py", "export.py",
                   "record.py", "run.py", "util.py", "video.py"},
}


def check_structure() -> None:
    for src in ("01_lingzao", "02_getnote"):
        d = ROOT / "01_采集" / src
        if not d.is_dir():
            add(ERROR, "采集层少了源文件夹", rel(d))
            continue
        miss = EXPECT_COLLECT - {f.name for f in d.iterdir() if f.is_file()}
        if miss:
            add(WARN, f"采集层缺文件（{src}）", "、".join(sorted(miss)))
        else:
            add(INFO, f"采集层完整（{src}）", "、".join(sorted(EXPECT_COLLECT)))

    for src, expect in EXPECT_PROCESS.items():
        d = ROOT / "03_加工" / src
        if not d.is_dir():
            add(ERROR, "加工层少了源文件夹", rel(d))
            continue
        miss = expect - {f.name for f in d.iterdir() if f.is_file()}
        if miss:
            add(WARN, f"加工层缺文件（{src}）", "、".join(sorted(miss)))
        else:
            add(INFO, f"加工层完整（{src}）", f"{len(expect)} 个 .py")

    for f in ("README.md", "05_技能/新增渠道规范.md"):
        if not (ROOT / f).is_file():
            add(WARN, "文档缺失", f)


# ---------------------------------------------------------------------------
# 7. 配置残留
# ---------------------------------------------------------------------------

STALE_NAMES = [
    "adapters/", "03_加工/run.py", "03_加工/scripts",
    "笔记信息.md", "kb blogger-", "collect.py cover",
]


def check_config() -> None:
    for d in sorted((ROOT / "01_采集").iterdir()):
        if not d.is_dir():
            continue
        for f in (d / "collect.py", d / "config.json", d / "README.md"):
            if not f.is_file():
                continue
            text = f.read_text(encoding="utf-8", errors="ignore")
            hit = [s for s in STALE_NAMES if s in text]
            if hit:
                add(INFO, f"提到了可能过期的名字（{rel(f)}）",
                    "、".join(hit) + "   ← 若是「说明它已废」就没事，自己看一眼")


# ---------------------------------------------------------------------------
# 8. 表里的「未知」
# ---------------------------------------------------------------------------

def check_unknown() -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        add(INFO, "跳过表格检查", "没装 openpyxl（pip install openpyxl）")
        return

    pack = PRODUCE / "博主"
    tables = []
    if pack.is_dir():
        tables += list(pack.glob("*/_索引.xlsx"))
    tables += list((PRODUCE / "表格").glob("*.xlsx")) if (PRODUCE / "表格").is_dir() else []

    if not tables:
        return
    for t in tables[:6]:
        try:
            ws = load_workbook(t, read_only=True).worksheets[0]
        except Exception:
            continue
        rows = list(ws.iter_rows(values_only=True))
        head = next((i for i, r in enumerate(rows[:15])
                     if r and any(str(c) in ("标题", "笔记ID") for c in r if c)), None)
        if head is None:
            add(WARN, "表格找不到表头行", rel(t))
            continue
        cols = [str(c) for c in rows[head] if c]
        body = [r for r in rows[head + 1:] if any(r)]
        unknown = sum(1 for r in body for c in r if str(c).strip() == "未知")
        zero_like = sum(1 for r in body for c in r if str(c).strip() in ("0", "0.0"))
        if unknown:
            add(INFO, "表里有「未知」单元格 —— 别当 0 相加",
                f"{t.name}：{len(body)} 行 × {len(cols)} 列，"
                f"「未知」{unknown} 处（另有真 0 共 {zero_like} 处）")
        else:
            add(INFO, "表里没有「未知」", f"{t.name}：{len(body)} 行")


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------

def walk_files(base: Path):
    for p in sorted(base.rglob("*")):
        if any(part in SKIP_DIRS for part in p.parts):
            continue
        if p.is_file():
            yield p


def main() -> int:
    quiet = "--quiet" in sys.argv

    print("项目自检")
    print("=" * 60)

    check_secrets()
    check_naming()
    check_subtitles()
    check_covers()
    check_reconcile()
    check_structure()
    check_config()
    check_unknown()

    n_err = sum(1 for x in _found if x[0] == ERROR)
    n_warn = sum(1 for x in _found if x[0] == WARN)
    n_info = sum(1 for x in _found if x[0] == INFO)

    for level in (ERROR, WARN):
        for lv, title, detail in _found:
            if lv != level:
                continue
            print(f"\n[{level}] {title}")
            if detail:
                print(f"        {detail}")

    if not quiet:
        print(f"\n{'-' * 60}")
        for lv, title, detail in _found:
            if lv != INFO:
                continue
            print(f"[{INFO}] {title}")
            if detail:
                print(f"        {detail}")

    print("\n" + "=" * 60)
    print(f"错误 {n_err} ｜ 警告 {n_warn} ｜ 提示 {n_info}"
          + ("" if quiet else "     （加 --quiet 只看前两类）"))

    if not STORE.is_dir() or not any(STORE.iterdir()):
        print("\n还没抓过数据 —— 这份自检大部分项目是空的。")
        print("  先跑一次采集，再看它。")

    if n_err:
        print("\n有「错误」级问题，动手前先处理。")
        return 1
    print("\n没有错误级问题。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
